"""Convert every supported input (PDF, DOCX, images, URL, text) into plain text.

All public functions either return non-empty text or raise ExtractionError with a
message that is safe to show directly to the user.
"""

import io
import ipaddress
import re
import socket
from functools import lru_cache
from urllib.parse import urljoin, urlparse

import numpy as np
import requests
from PIL import Image, ImageOps, UnidentifiedImageError

MAX_FILE_BYTES = 15 * 1024 * 1024  # 15 MB per uploaded file
MAX_IMAGES = 20
MAX_PDF_PAGES = 60  # keeps OCR of scanned PDFs to a few minutes at most
MAX_TEXT_CHARS = 200_000
MAX_URL_BYTES = 5 * 1024 * 1024
URL_TIMEOUT = 10  # seconds
MAX_REDIRECTS = 5
MIN_PAGE_CHARS = 40  # PDF pages with less text than this are treated as scanned and OCR'd
OCR_MAX_SIDE = 2000  # large photos are downscaled before OCR for speed
OCR_MIN_WIDTH = 1600  # narrower images are upscaled (up to 2x) before OCR
MERGED_WORD_MIN = 10  # letter runs at least this long are checked for missing spaces


class ExtractionError(Exception):
    """Raised when an input cannot be turned into text. The message is user-facing."""


def _check_size(data: bytes, what: str) -> None:
    if not data:
        raise ExtractionError(f"The {what} is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ExtractionError(f"The {what} is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB.")


def _finish(text: str, empty_message: str) -> str:
    text = "\n".join(line.rstrip() for line in text.replace("\x00", "").splitlines()).strip()
    if not text:
        raise ExtractionError(empty_message)
    return text[:MAX_TEXT_CHARS]


# ---------- OCR ----------

@lru_cache(maxsize=1)
def _ocr_engine():
    from rapidocr_onnxruntime import RapidOCR  # imported lazily: loading takes a second

    return RapidOCR()


_LETTER_RUN = re.compile(r"[A-Za-z]{%d,}" % MERGED_WORD_MIN)


def _fix_merged_words(line: str) -> str:
    """OCR on small or blurry photos often drops spaces ("theTermsofService").
    Long letter-only runs are re-split with a word-frequency dictionary; real long words stay intact."""
    import wordninja

    def split(match: re.Match) -> str:
        parts = wordninja.split(match.group(0))
        return " ".join(parts) if parts else match.group(0)

    return _LETTER_RUN.sub(split, line)


def _ocr_image(img: Image.Image) -> str:
    img = ImageOps.exif_transpose(img).convert("RGB")  # respect phone camera rotation
    if max(img.size) > OCR_MAX_SIDE:
        img.thumbnail((OCR_MAX_SIDE, OCR_MAX_SIDE))
    elif img.width < OCR_MIN_WIDTH:  # small text: upscale so the recogniser sees word gaps
        scale = min(2.0, OCR_MIN_WIDTH / img.width)
        img = img.resize((int(img.width * scale), int(img.height * scale)), Image.LANCZOS)
    result, _ = _ocr_engine()(np.array(img))
    if not result:
        return ""
    # result rows: [box, text, score]; group boxes into lines by their vertical position.
    rows = sorted(result, key=lambda r: (r[0][0][1], r[0][0][0]))
    lines: list[list] = []
    for box, text, _score in rows:
        top, height = box[0][1], abs(box[3][1] - box[0][1]) or 1
        if lines and abs(top - lines[-1][0]) < height * 0.6:
            lines[-1][1].append((box[0][0], text))
        else:
            lines.append([top, [(box[0][0], text)]])
    return "\n".join(_fix_merged_words(" ".join(t for _, t in sorted(words))) for _, words in lines)


# ---------- public extractors ----------

def from_text(text: str) -> str:
    return _finish(text or "", "Please paste some text to analyze.")


def from_pdf(data: bytes) -> str:
    import pymupdf

    _check_size(data, "PDF file")
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ExtractionError("This file could not be opened as a PDF.") from exc
    if doc.needs_pass:
        raise ExtractionError("This PDF is password-protected. Please upload an unlocked copy.")
    if doc.page_count > MAX_PDF_PAGES:
        raise ExtractionError(f"This PDF has {doc.page_count} pages. Please upload at most {MAX_PDF_PAGES} pages.")

    pages = []
    for page in doc:
        text = page.get_text("text")
        if len(text.strip()) < MIN_PAGE_CHARS:  # scanned page: render and OCR it
            pix = page.get_pixmap(dpi=200)
            text = _ocr_image(Image.open(io.BytesIO(pix.tobytes("png"))))
        pages.append(text)
    return _finish("\n\n".join(pages), "No readable text was found in this PDF.")


def from_docx(data: bytes) -> str:
    from docx import Document

    _check_size(data, "Word file")
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("This file could not be opened as a Word (.docx) document.") from exc

    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:  # merged cells repeat, so de-duplicate per row
                if cell.text and cell.text not in seen:
                    seen.add(cell.text)
                    parts.append(cell.text)
    return _finish("\n".join(parts), "No text was found in this Word document.")


def from_images(images: list[bytes]) -> str:
    if not images:
        raise ExtractionError("Please upload at least one image.")
    if len(images) > MAX_IMAGES:
        raise ExtractionError(f"Please upload at most {MAX_IMAGES} images at a time.")
    pages = []
    for i, data in enumerate(images, start=1):
        _check_size(data, f"image {i}")
        try:
            img = Image.open(io.BytesIO(data))
            img.load()
        except Image.DecompressionBombError as exc:
            raise ExtractionError(f"Image {i} has too many pixels. Use a normal photo or screenshot.") from exc
        except (UnidentifiedImageError, OSError) as exc:
            raise ExtractionError(f"Image {i} could not be read. Use JPG or PNG.") from exc
        pages.append(_ocr_image(img))
    return _finish(
        "\n\n".join(p for p in pages if p.strip()),
        "No text could be read from the image(s). Try a sharper, well-lit photo taken straight on.",
    )


def _ensure_public_http_url(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise ExtractionError("Please enter a full web link starting with http:// or https://")
    try:
        infos = socket.getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80))
    except socket.gaierror as exc:
        raise ExtractionError("That website address could not be found. Check the link.") from exc
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if not ip.is_global:
            raise ExtractionError("Links to local or private network addresses are not allowed.")
    return url


def _decode_html(body: bytes, declared: str | None) -> str:
    """Prefer UTF-8 (most of the web); fall back to the declared charset only if UTF-8 fails.
    requests declares ISO-8859-1 whenever the server omits a charset, which garbles UTF-8 pages."""
    try:
        return body.decode("utf-8")
    except UnicodeDecodeError:
        return body.decode(declared or "latin-1", errors="replace")


def from_url(url: str) -> str:
    import trafilatura

    url = (url or "").strip()
    if url and "://" not in url:
        url = "https://" + url

    headers = {"User-Agent": "Mozilla/5.0 (ClauseGuard student project; terms analyser)"}
    try:
        # Follow redirects manually so every hop is checked against private addresses.
        for _ in range(MAX_REDIRECTS + 1):
            _ensure_public_http_url(url)
            resp = requests.get(url, headers=headers, timeout=URL_TIMEOUT, stream=True, allow_redirects=False)
            if resp.is_redirect:
                url = urljoin(url, resp.headers.get("Location", ""))
                resp.close()
                continue
            break
        else:
            raise ExtractionError("The link redirected too many times.")

        if resp.status_code in (401, 403, 429):
            raise ExtractionError("This website blocked automatic reading. Copy the text and use 'Paste text' instead.")
        if resp.status_code >= 400:
            raise ExtractionError(f"The website returned an error (HTTP {resp.status_code}). Check the link, or copy the text and use 'Paste text'.")

        body = b""
        for chunk in resp.iter_content(64 * 1024):
            body += chunk
            if len(body) > MAX_URL_BYTES:
                raise ExtractionError("That page is too large to analyze. Copy the terms section and use 'Paste text' instead.")
    except requests.RequestException as exc:
        raise ExtractionError("Could not reach that website. Check the link and your internet connection.") from exc

    content_type = resp.headers.get("Content-Type", "").lower()
    if "pdf" in content_type or url.lower().endswith(".pdf"):
        return from_pdf(body)

    html = _decode_html(body, resp.encoding)
    text = trafilatura.extract(html, include_comments=False, include_tables=True, favor_recall=True) or ""
    return _finish(
        text,
        "No terms text could be found on that page (it may need JavaScript). "
        "Copy the text and use 'Paste text' instead.",
    )
