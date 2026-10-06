import io

import pymupdf
import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont

from src import extract
from src.extract import ExtractionError

SENTENCE = "We may terminate your account at any time without notice."


def _font(size=34):
    for name in ("arial.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def _text_image(text: str) -> bytes:
    img = Image.new("RGB", (1400, 160), "white")
    ImageDraw.Draw(img).text((30, 50), text, fill="black", font=_font())
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _text_pdf(text: str) -> bytes:
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), text, fontsize=11)
    return doc.tobytes()


def _scanned_pdf(text: str) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page(width=1400, height=400)
    page.insert_image(pymupdf.Rect(0, 0, 1400, 160), stream=_text_image(text))
    return doc.tobytes()


def test_text_pdf():
    assert SENTENCE in extract.from_pdf(_text_pdf(SENTENCE))


def test_scanned_pdf_uses_ocr():
    text = extract.from_pdf(_scanned_pdf(SENTENCE)).lower()
    assert "terminate your account" in text


def test_docx_with_table():
    doc = Document()
    doc.add_paragraph(SENTENCE)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "The deposit is non-refundable."
    buf = io.BytesIO()
    doc.save(buf)
    text = extract.from_docx(buf.getvalue())
    assert SENTENCE in text
    assert "The deposit is non-refundable." in text


def test_images_are_joined_in_order():
    text = extract.from_images([_text_image("First page terms apply."), _text_image("Second page fees apply.")])
    assert text.lower().index("first page") < text.lower().index("second page")


def test_blank_image_raises():
    buf = io.BytesIO()
    Image.new("RGB", (400, 200), "white").save(buf, format="PNG")
    with pytest.raises(ExtractionError):
        extract.from_images([buf.getvalue()])


def test_invalid_pdf_raises():
    with pytest.raises(ExtractionError):
        extract.from_pdf(b"not a pdf")


@pytest.mark.parametrize(
    "url",
    ["http://localhost/terms", "http://127.0.0.1/terms", "http://192.168.1.10/terms", "file:///etc/passwd", "ftp://x.com/a"],
)
def test_url_blocks_private_hosts(url):
    with pytest.raises(ExtractionError):
        extract.from_url(url)


def test_empty_text_raises():
    with pytest.raises(ExtractionError):
        extract.from_text("   \n ")


def test_text_cap():
    text = extract.from_text("a" * (extract.MAX_TEXT_CHARS + 500))
    assert len(text) == extract.MAX_TEXT_CHARS


def test_oversized_file_rejected():
    with pytest.raises(ExtractionError):
        extract.from_pdf(b"0" * (extract.MAX_FILE_BYTES + 1))


def test_html_without_charset_is_decoded_as_utf8():
    # requests reports ISO-8859-1 when a server omits the charset; UTF-8 pages must still decode correctly.
    assert extract._decode_html("Dropbox’s terms".encode("utf-8"), "ISO-8859-1") == "Dropbox’s terms"


def test_html_with_real_latin1_still_works():
    assert extract._decode_html("café".encode("latin-1"), "ISO-8859-1") == "café"


def test_merged_ocr_words_are_split_and_punctuation_kept():
    assert extract._fix_merged_words("Lastupdated:1August2026, theTermsofService.") == (
        "Last updated:1August2026, the Terms of Service."
    )


def test_normal_words_untouched():
    line = "Your responsibilities regarding confidentiality remain."
    assert extract._fix_merged_words(line) == line


def test_decompression_bomb_image_gives_friendly_error(monkeypatch):
    monkeypatch.setattr(Image, "MAX_IMAGE_PIXELS", 1000)  # make a normal image count as a "bomb"
    with pytest.raises(ExtractionError):
        extract.from_images([_text_image("tiny")])


def test_pdf_page_limit():
    doc = pymupdf.open()
    for _ in range(extract.MAX_PDF_PAGES + 1):
        doc.new_page().insert_text((72, 72), SENTENCE)
    with pytest.raises(ExtractionError):
        extract.from_pdf(doc.tobytes())
