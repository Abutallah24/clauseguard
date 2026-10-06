"""Build the 2-3 page Word report from the real results in reports/metrics.json.

Run:  .venv\\Scripts\\python -m reports.build_report
Output: reports/ClauseGuard_Report.docx
"""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from src.labels import LABELS

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
FIG = REPORTS / "figures"
SHOTS = REPORTS / "screenshots"
OUT = REPORTS / "ClauseGuard_Report.docx"

INK = RGBColor(0x0D, 0x25, 0x3D)
INK_2 = RGBColor(0x3C, 0x4F, 0x66)
ACCENT = RGBColor(0x3A, 0x3F, 0xD1)
ORDER = ["tfidf_logreg", "tfidf_svm", "distilbert"]


# ---------- helpers ----------

def set_cell_background(cell, hex_fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_table_borders(table, color="C3CBD6", inner=True) -> None:
    borders = OxmlElement("w:tblBorders")
    edges = ["top", "bottom"] + (["insideH"] if inner else [])
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{edge}")
        if edge in edges:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    table._tbl.tblPr.append(borders)


def para(doc, text="", size=10.5, bold=False, color=INK, align=None, space_after=4, italic=False):
    p = doc.add_paragraph()
    if text:
        run = p.add_run(text)
        run.font.size, run.bold, run.italic = Pt(size), bold, italic
        run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    return p


def rich(doc, parts, size=10.5, space_after=4):
    """parts: list of (text, bold) tuples in one paragraph."""
    p = doc.add_paragraph()
    for text, bold in parts:
        run = p.add_run(text)
        run.font.size, run.bold = Pt(size), bold
        run.font.color.rgb = INK
    p.paragraph_format.space_after = Pt(space_after)
    return p


def heading(doc, text):
    p = doc.add_heading(text, level=2)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    for run in p.runs:
        run.font.size, run.font.color.rgb = Pt(12.5), ACCENT
    return p


def bullet(doc, parts, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    for text, bold in parts:
        run = p.add_run(text)
        run.font.size, run.bold = Pt(size), bold
        run.font.color.rgb = INK
    p.paragraph_format.space_after = Pt(1.5)
    return p


def figure(doc, path: Path, width_cm: float, caption: str):
    if not path.exists():
        return
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_after = Pt(1)
    para(doc, caption, size=8.5, color=INK_2, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6, italic=True)


def side_by_side(doc, items):
    """items: list of (path, width_cm, caption) placed in one borderless row; missing images are skipped."""
    items = [item for item in items if item[0].exists()]
    if not items:
        return
    table = doc.add_table(rows=2, cols=len(items))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, inner=False)
    for i, (path, width, caption) in enumerate(items):
        cell = table.cell(0, i)
        if path.exists():
            cell.paragraphs[0].add_run().add_picture(str(path), width=Cm(width))
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = table.cell(1, i).paragraphs[0]
        run = cap.add_run(caption)
        run.font.size, run.italic = Pt(8.5), True
        run.font.color.rgb = INK_2
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "", space_after=2)


# ---------- content ----------

def build() -> None:
    metrics = json.loads((REPORTS / "metrics.json").read_text())
    eda = json.loads((REPORTS / "eda.json").read_text())
    models = metrics["models"]
    keys = [k for k in ORDER if k in models]
    best_key = metrics["best_model"]
    best = models[best_key]
    sizes = eda["split_sizes"]
    total = sum(sizes.values())

    doc = Document()
    for section in doc.sections:
        section.page_height, section.page_width = Cm(29.7), Cm(21.0)
        section.top_margin = section.bottom_margin = Cm(1.6)
        section.left_margin = section.right_margin = Cm(1.9)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.line_spacing = 1.08

    # Title block
    para(doc, "K. J. Somaiya School of Engineering, Mumbai · Somaiya Vidyavihar University", 9, color=INK_2,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    para(doc, "Artificial Intelligence (216U41C501) · Lab CA Mini-Project · TY CCE (EXTC) · Sem V · 2025-26", 9,
         color=INK_2, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    para(doc, "ClauseGuard: Detecting Unfair Clauses in Terms & Conditions with NLP", 17, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    para(doc, "Name: ____________________   Roll No.: ________   Division: ____   "
              "Faculty: Mr. Ankit Khivasara", 10, color=INK_2, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    heading(doc, "Abstract")
    para(doc, (
        f"People routinely accept terms and conditions without reading them. ClauseGuard is an NLP system that reads a "
        f"document and flags clauses that are potentially unfair to the user. We trained and compared three classifiers "
        f"on the UNFAIR-ToS dataset ({total:,} sentences from 50 online terms of service, 8 unfair-clause types): "
        f"TF-IDF with Logistic Regression, TF-IDF with a Linear SVM, and a fine-tuned DistilBERT transformer. "
        f"On the held-out test set the best model ({best['name']}) reached a macro F1 of {best['macro']['f1']:.3f} "
        f"and an F1 of {best['binary']['f1']:.3f} for separating fair from unfair clauses. The model is deployed in a "
        f"web app that accepts PDFs, scanned documents, Word files, photos, camera shots, links and pasted text, "
        f"highlights risky clauses with plain-English explanations and computes a 0–100 risk score."
    ))

    heading(doc, "1. Problem statement")
    para(doc, (
        "Terms of service, rental agreements and offer letters are long and written in legal language, so most people "
        "agree without reading them. A Deloitte survey of 2,000 US consumers found that 91% accept legal terms without "
        "reading them (97% of 18–34 year olds) [4]. In a study by Obar and Oeldorf-Hirsch, 98% of students missed a "
        "clause that gave away their first-born child [5], and reading every privacy policy a person meets in a year "
        "would take about 76 working days [6]. Users therefore unknowingly accept arbitration instead of courts, "
        "one-sided termination and changes, liability waivers, auto-renewals and penalties. The goal is a tool that "
        "automatically points out such clauses and quantifies the overall risk before the user agrees."
    ))

    heading(doc, "2. Dataset")
    unfair_pct = eda["unfair_share"]["train"]
    para(doc, (
        f"We used UNFAIR-ToS from the LexGLUE benchmark [2], built on the CLAUDETTE corpus [1]. It contains {total:,} "
        f"sentences from 50 terms of service (e.g. of large online platforms), labelled by legal experts with zero or "
        f"more of 8 unfair types: limitation of liability, unilateral termination, unilateral change, content removal, "
        f"contract by using, choice of law, jurisdiction and arbitration. We kept the official splits: "
        f"{sizes['train']:,} training, {sizes['validation']:,} validation and {sizes['test']:,} test sentences. "
        f"Only {unfair_pct:.1%} of training sentences are unfair, so the data is strongly imbalanced; "
        f"{eda['multi_label_sentences']} sentences carry two labels. We framed the task as 9-class classification "
        f"(Fair + 8 types); multi-label training sentences were duplicated once per label."
    ))

    heading(doc, "3. Methodology")
    figure(doc, FIG / "architecture.png", 17.0, "Figure 1. ClauseGuard pipeline.")
    bullet(doc, [("Text extraction: ", True), ("PyMuPDF for PDFs; pages with no text layer are rendered and read with "
                 "RapidOCR (also used for photos and the camera, with up-scaling and a word-segmentation fix for "
                 "merged words); python-docx for Word files; trafilatura for web pages.", False)])
    bullet(doc, [("Clause splitting: ", True), ("a rule-based splitter on sentence ends and numbered items that drops "
                 "headings.", False)])
    bullet(doc, [("Classical models: ", True), ("TF-IDF features (word 1–2-grams + character 3–5-grams, sub-linear TF) "
                 "with Logistic Regression and with a Linear SVM (probabilities via Platt-style calibration). Both use "
                 "class_weight = balanced to counter the imbalance.", False)])
    if "distilbert" in models:
        hp = models["distilbert"].get("hyperparameters", {})
        bullet(doc, [("Deep learning: ", True), (f"DistilBERT (distilbert-base-uncased) fine-tuned for "
                     f"{hp.get('epochs', 3)} epochs on CPU (AdamW, lr {hp.get('learning_rate', 5e-5)}, batch "
                     f"{hp.get('batch_size', 16)}, max length {hp.get('max_len', 128)}, linear warm-up) with a "
                     "cross-entropy loss weighted by the square root of balanced class weights (milder than full balancing, which made training unstable on rare classes); the epoch with the best validation macro F1 was kept.", False)])
    bullet(doc, [("Rule checks: ", True), ("5 transparent keyword rules for patterns outside the dataset that "
                 "students often meet (auto-renewal, hidden fees, data sharing, bond / lock-in, deposit forfeiture), "
                 "always labelled separately from model predictions.", False)])
    bullet(doc, [("Risk score: ", True), ("each finding contributes risk = severity of its type × model confidence; "
                 "score = 100 × (1 − e^(−Σ risk / 2.5)), giving Low (0–29), Medium (30–60) or High (61–100).", False)])

    heading(doc, "4. Results")
    para(doc, "All metrics are on the official test split, which was not used for training or model selection.",
         space_after=3)
    header = ["Model", "Accuracy", "Macro P", "Macro R", "Macro F1", "Unfair-only macro F1", "Fair vs unfair F1",
              "Train time"]
    table = doc.add_table(rows=1 + len(keys), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    widths = [4.6, 1.6, 1.5, 1.5, 1.6, 2.2, 2.1, 1.7]
    for i, h in enumerate(header):
        cell = table.cell(0, i)
        cell.width = Cm(widths[i])
        run = cell.paragraphs[0].add_run(h)
        run.bold, run.font.size = True, Pt(8.5)
        set_cell_background(cell, "EEF0FF")
    for r, key in enumerate(keys, start=1):
        m = models[key]
        secs = m.get("train_seconds", 0)
        values = [m["name"] + (" (best)" if key == best_key else ""), f"{m['accuracy']:.3f}",
                  f"{m['macro']['precision']:.3f}", f"{m['macro']['recall']:.3f}", f"{m['macro']['f1']:.3f}",
                  f"{m['unfair_only']['macro_f1']:.3f}", f"{m['binary']['f1']:.3f}",
                  f"{secs / 60:.0f} min" if secs >= 60 else f"{secs:.0f} s"]
        for c, v in enumerate(values):
            cell = table.cell(r, c)
            cell.width = Cm(widths[c])
            run = cell.paragraphs[0].add_run(v)
            run.font.size, run.bold = Pt(8.5), key == best_key
            if c:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para(doc, "", space_after=3)

    side_by_side(doc, [
        (FIG / "model_comparison.png", 8.6, "Figure 2. Model comparison (test split)."),
        (FIG / f"confusion_{best_key}.png", 7.6, f"Figure 3. Confusion matrix, {best['name']}."),
    ])

    per_class = best["per_class"]
    unfair = {k: v for k, v in per_class.items() if k not in ("Fair", "accuracy", "macro avg", "weighted avg")
              and isinstance(v, dict)}
    strongest = max(unfair, key=lambda k: unfair[k]["f1-score"])
    weakest = min(unfair, key=lambda k: unfair[k]["f1-score"])
    test_best = max(keys, key=lambda k: models[k]["macro"]["f1"])
    fair_share = 1 - eda["unfair_share"]["test"]

    cm = best["confusion_matrix"]
    labels = LABELS
    false_alarms = sum(cm[0][1:])  # fair predicted as unfair
    missed = sum(row[0] for row in cm[1:])  # unfair predicted as fair
    pairs = [(cm[i][j], i, j) for i in range(1, len(cm)) for j in range(1, len(cm)) if i != j]
    top_count, top_i, top_j = max(pairs)
    selection = (f"{best['name']} was selected for the app because it has the best validation macro F1 "
                 f"({best['validation_macro_f1']:.3f}).")
    if test_best != best_key:
        selection += (f" On the test set {models[test_best]['name']} scores slightly higher "
                      f"({models[test_best]['macro']['f1']:.3f} vs {best['macro']['f1']:.3f}), so the two are close.")
    errors = (f"For {best['name']}, {false_alarms} fair sentences were flagged as unfair (false alarms) and "
              f"{missed} unfair sentences were missed.")
    if top_count:
        errors += (f" The most common mix-up between unfair types was {labels[top_i]} predicted as "
                   f"{labels[top_j]} ({top_count} sentences).")
    para(doc, (
        f"Accuracy is high for every model because about {fair_share:.0%} of test sentences are fair, so macro F1 "
        f"(all 9 classes weighted equally) is the main metric. {selection} Per class, it is strongest on {strongest} "
        f"(F1 {unfair[strongest]['f1-score']:.2f}) and weakest on {weakest} (F1 {unfair[weakest]['f1-score']:.2f}, "
        f"{int(unfair[weakest]['support'])} test sentences). {errors}"
    ))

    heading(doc, "5. Application")
    figure(doc, SHOTS / "report_desktop.png", 16.5,
           "Figure 4. ClauseGuard web app: risk score, category summary and the document with risky clauses highlighted.")
    para(doc, (
        "The Streamlit app offers tabs for PDF, Word, multiple images, camera, link and pasted text, three built-in "
        "sample documents (app terms, a rental agreement and an internship offer letter), a downloadable HTML report and "
        "a model-performance page. It runs on a laptop and on phones over the same Wi-Fi. All seven input types were "
        "tested end-to-end with automated browser tests, alongside a pytest suite for extraction, clause splitting, "
        "rules and scoring."
    ))

    heading(doc, "6. Conclusion and future scope")
    para(doc, (
        "ClauseGuard shows that a model trained on a modest labelled dataset can reliably flag unfair clauses and make "
        "long legal documents understandable in seconds. Its main limitation is domain: it learned from English online "
        "terms of service, so it is less accurate on Indian rental agreements and offer letters, where the rule checks "
        "help. Future work: collect and label Indian contracts (rent, employment, loan apps), support Hindi and Marathi, "
        "fine-tune a legal-domain model such as LegalBERT, suggest fairer rewordings of risky clauses, and ship a "
        "browser extension that scans terms before the user clicks \"I agree\". ClauseGuard is not legal advice."
    ))

    heading(doc, "References")
    refs = [
        "M. Lippi et al., \"CLAUDETTE: an automated detector of potentially unfair clauses in online terms of service,\" "
        "Artificial Intelligence and Law, vol. 27, pp. 117–139, 2019.",
        "I. Chalkidis et al., \"LexGLUE: A Benchmark Dataset for Legal Language Understanding in English,\" in Proc. ACL, 2022.",
        "V. Sanh, L. Debut, J. Chaumond and T. Wolf, \"DistilBERT, a distilled version of BERT: smaller, faster, cheaper "
        "and lighter,\" arXiv:1910.01108, 2019.",
        "Deloitte, 2017 Global Mobile Consumer Survey (US edition), as reported by Business Insider, Nov. 2017.",
        "J. A. Obar and A. Oeldorf-Hirsch, \"The biggest lie on the Internet: ignoring the privacy policies and terms of "
        "service policies of social networking services,\" Information, Communication & Society, 23(1), 2020.",
        "A. M. McDonald and L. F. Cranor, \"The Cost of Reading Privacy Policies,\" I/S: A Journal of Law and Policy for "
        "the Information Society, 4(3), 2008.",
    ]
    for i, ref in enumerate(refs, start=1):
        para(doc, f"[{i}] {ref}", size=8.5, color=INK_2, space_after=1)

    doc.save(OUT)
    print("report written to", OUT)


if __name__ == "__main__":
    build()
