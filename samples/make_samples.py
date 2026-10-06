"""Generate real test files for every input type from sample_texts.py.

Creates in samples/:  *.txt, *.pdf (text PDF), *.docx, rental_agreement_scanned.pdf (image-only PDF),
internship_offer_page{1,2}.png (photos of pages) and camera_photo.jpg (tilted, phone-like photo).
Run:  .venv\\Scripts\\python -m samples.make_samples
"""

import io
import textwrap
from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFilter, ImageFont

from samples.sample_texts import SAMPLES

OUT = Path(__file__).resolve().parent
PAGE_W, PAGE_H, MARGIN = 1240, 1754, 90  # A4 at 150 dpi


def _font(size: int, bold: bool = False):
    for name in (("arialbd.ttf" if bold else "arial.ttf"), "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def write_pdf(text: str, path: Path) -> None:
    doc = pymupdf.open()
    rect = pymupdf.Rect(56, 56, 539, 786)
    remaining = text
    while remaining:
        page = doc.new_page()
        # insert_textbox returns <0 when text overflows; split by paragraphs page by page.
        paras, chunk = remaining.split("\n"), []
        while paras:
            trial = "\n".join(chunk + [paras[0]])
            probe = pymupdf.open().new_page()
            if probe.insert_textbox(rect, trial, fontsize=10.5, fontname="helv") < 0 and chunk:
                break
            chunk.append(paras.pop(0))
        page.insert_textbox(rect, "\n".join(chunk), fontsize=10.5, fontname="helv")
        remaining = "\n".join(paras)
    doc.save(path)


def write_docx(title: str, text: str, path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(11)
    lines = text.strip().splitlines()
    doc.add_heading(lines[0], level=1)
    for line in lines[1:]:
        if line.strip():
            doc.add_paragraph(line.strip())
    # A small table, so the DOCX extractor's table support is exercised too.
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text, table.cell(0, 1).text = "Document", title
    table.cell(1, 0).text, table.cell(1, 1).text = "Note", "Generated sample for ClauseGuard demo"
    doc.save(path)


def render_pages(text: str) -> list[Image.Image]:
    """Render text onto white A4 page images (like a scan)."""
    font, line_h = _font(26), 38
    max_chars = 78
    lines: list[str] = []
    for para in text.strip().splitlines():
        lines.extend(textwrap.wrap(para, max_chars) or [""])
    per_page = (PAGE_H - 2 * MARGIN) // line_h
    pages = []
    for start in range(0, len(lines), per_page):
        img = Image.new("RGB", (PAGE_W, PAGE_H), "white")
        draw = ImageDraw.Draw(img)
        for i, line in enumerate(lines[start:start + per_page]):
            draw.text((MARGIN, MARGIN + i * line_h), line, fill=(25, 25, 25), font=font)
        pages.append(img)
    return pages


def write_scanned_pdf(text: str, path: Path) -> None:
    doc = pymupdf.open()
    for img in render_pages(text):
        buf = io.BytesIO()
        img.filter(ImageFilter.GaussianBlur(0.4)).save(buf, format="PNG")
        page = doc.new_page(width=595, height=842)
        page.insert_image(page.rect, stream=buf.getvalue())
    doc.save(path)


def camera_photo(page: Image.Image) -> Image.Image:
    """Imitate a phone photo: slight rotation, warm background, mild blur."""
    rotated = page.rotate(1.5, expand=True, fillcolor=(205, 190, 165), resample=Image.BICUBIC)
    canvas = Image.new("RGB", (rotated.width + 160, rotated.height + 160), (190, 175, 150))
    canvas.paste(rotated, (80, 80))
    return canvas.filter(ImageFilter.GaussianBlur(0.6)).resize((canvas.width * 3 // 4, canvas.height * 3 // 4))


def main() -> None:
    for key, sample in SAMPLES.items():
        (OUT / f"{key}.txt").write_text(sample["text"], encoding="utf-8")
        write_pdf(sample["text"], OUT / f"{key}.pdf")
        write_docx(sample["title"], sample["text"], OUT / f"{key}.docx")

    write_scanned_pdf(SAMPLES["rental_agreement"]["text"], OUT / "rental_agreement_scanned.pdf")

    pages = render_pages(SAMPLES["internship_offer"]["text"])
    half = len(SAMPLES["internship_offer"]["text"].strip().splitlines()) // 2
    lines = SAMPLES["internship_offer"]["text"].strip().splitlines()
    if len(pages) == 1:  # split into two pages so the multi-image upload can be demoed
        pages = render_pages("\n".join(lines[:half])) + render_pages("\n".join(lines[half:]))
    for i, img in enumerate(pages[:2], start=1):
        img.save(OUT / f"internship_offer_page{i}.png")

    camera_photo(render_pages(SAMPLES["streamly_terms"]["text"])[0]).save(OUT / "camera_photo.jpg", quality=88)
    print("samples written:", sorted(p.name for p in OUT.iterdir() if p.suffix in {".pdf", ".docx", ".png", ".jpg", ".txt"}))


if __name__ == "__main__":
    main()
